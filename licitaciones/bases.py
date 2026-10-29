"""Extrae el texto de las bases (PDF/Word) y deja solo las partes que importan para ofertar.

Las bases suelen tener 30-80 páginas de texto legal repetido. Mandarlas enteras a Claude gasta
decenas de miles de tokens; aquí se quedan los párrafos sobre requisitos, evaluación, anexos,
plazos, garantías, etc., con un tope de largo.
"""
import re
from pathlib import Path

from .filtro import normalizar

CLAVES = [
    "criterio", "evaluaci", "ponderaci", "puntaje", "requisito", "admisib", "inadmisib", "anexo",
    "formulario", "declaraci", "garantia", "boleta", "plazo", "fecha", "documento", "antecedente",
    "oferta tecnica", "oferta economica", "propuesta tecnica", "experiencia", "relator", "facilitador",
    "metodolog", "contenido", "objetivo", "producto", "entregable", "informe", "participantes",
    "horas", "sesiones", "duracion", "lugar", "modalidad", "presupuesto", "monto", "precio", "valor",
    "multa", "sence", "otec", "certific", "curriculum", "perfil", "consulta", "visita",
]
PATRON = re.compile(r"\b(?:" + "|".join(re.escape(c) for c in CLAVES) + ")")
TITULO = re.compile(r"^\s*(?:\d+(?:\.\d+)*[.)]?|[IVXL]+[.)]|[A-ZÁÉÍÓÚÑ][A-ZÁÉÍÓÚÑ ,]{6,})\s*\S")


def texto_de(ruta):
    ext = ruta.suffix.lower()
    if ext == ".pdf":
        from pypdf import PdfReader
        return "\n".join((p.extract_text() or "") for p in PdfReader(str(ruta)).pages)
    if ext == ".docx":
        from docx import Document
        doc = Document(str(ruta))
        partes = [p.text for p in doc.paragraphs]
        for tabla in doc.tables:
            for fila in tabla.rows:
                partes.append(" | ".join(c.text.strip() for c in fila.cells))
        return "\n".join(partes)
    if ext in (".txt", ".md"):
        return ruta.read_text(encoding="utf-8", errors="ignore")
    return ""


def _parrafos(texto):
    """Agrupa líneas en párrafos; un título corto queda pegado al párrafo que lo sigue."""
    bloques, actual = [], []
    for linea in texto.splitlines():
        linea = linea.strip()
        if not linea:
            if actual:
                bloques.append(" ".join(actual))
                actual = []
            continue
        if actual and TITULO.match(linea) and len(" ".join(actual)) > 200:
            bloques.append(" ".join(actual))
            actual = []
        actual.append(linea)
    if actual:
        bloques.append(" ".join(actual))
    return bloques


def recortar(texto, max_caracteres=30000):
    """Se queda con los párrafos que tocan temas clave (y el siguiente, que suele ser la continuación)."""
    parrafos = _parrafos(texto)
    elegidos, vistos = [], set()
    for i, p in enumerate(parrafos):
        if PATRON.search(normalizar(p)):
            for j in (i, i + 1):
                if j < len(parrafos) and j not in vistos and len(parrafos[j]) > 15:
                    vistos.add(j)
                    elegidos.append(j)
    salida, total = [], 0
    for j in sorted(elegidos):
        p = parrafos[j]
        if total + len(p) > max_caracteres:
            salida.append("[… recortado por largo: revisar bases_texto.txt si falta algo …]")
            break
        salida.append(p)
        total += len(p)
    return "\n\n".join(salida)


def procesar(carpeta_oferta, max_caracteres=30000):
    """Lee todo lo que haya en <oferta>/bases/ y escribe bases_texto.txt (completo) y bases_resumen.md (recortado)."""
    carpeta_bases = carpeta_oferta / "bases"
    archivos = sorted(p for p in carpeta_bases.glob("**/*") if p.is_file() and p.suffix.lower() in (".pdf", ".docx", ".txt", ".md")
                      and p.name != "LEEME.txt")
    if not archivos:
        return None
    completo, resumen = [], []
    por_archivo = max(4000, max_caracteres // len(archivos))
    for ruta in archivos:
        try:
            texto = texto_de(ruta)
        except Exception as e:  # PDF dañado o escaneado: se informa y se sigue con los demás
            texto = ""
            resumen.append(f"## {ruta.name}\n\nNo se pudo leer ({e}).")
            continue
        completo.append(f"===== {ruta.name} =====\n{texto}")
        if not texto.strip():
            resumen.append(f"## {ruta.name}\n\nSin texto extraíble (¿PDF escaneado?). Revisarlo a mano.")
            continue
        resumen.append(f"## {ruta.name}\n\n{recortar(texto, por_archivo)}")
    (carpeta_oferta / "bases_texto.txt").write_text("\n\n".join(completo), encoding="utf-8")
    texto_resumen = "# Extracto de las bases (solo partes clave)\n\n" + "\n\n".join(resumen) + "\n"
    (carpeta_oferta / "bases_resumen.md").write_text(texto_resumen, encoding="utf-8")
    total = sum(len(t) for t in completo)
    return {"archivos": len(archivos), "caracteres_total": total, "caracteres_resumen": len(texto_resumen)}
