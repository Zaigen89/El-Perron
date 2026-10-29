"""Convierte los borradores en Markdown (los que escribe Claude) a Word, listos para revisar y subir."""
import re

from docx import Document
from docx.shared import Pt

NEGRITA = re.compile(r"\*\*(.+?)\*\*")


def _texto(parrafo, texto):
    pos = 0
    for m in NEGRITA.finditer(texto):
        parrafo.add_run(texto[pos:m.start()])
        parrafo.add_run(m.group(1)).bold = True
        pos = m.end()
    parrafo.add_run(texto[pos:])


def _tabla(doc, filas):
    celdas = [[c.strip() for c in f.strip().strip("|").split("|")] for f in filas
              if not re.fullmatch(r"\|?[\s:|-]+\|?", f.strip())]
    if not celdas:
        return
    ancho = max(len(f) for f in celdas)
    tabla = doc.add_table(rows=len(celdas), cols=ancho)
    tabla.style = "Table Grid"
    for i, fila in enumerate(celdas):
        for j, valor in enumerate(fila):
            celda = tabla.cell(i, j)
            celda.text = ""
            _texto(celda.paragraphs[0], valor)
            if i == 0:
                for run in celda.paragraphs[0].runs:
                    run.bold = True


def md_a_docx(md, destino):
    doc = Document()
    doc.styles["Normal"].font.name = "Calibri"
    doc.styles["Normal"].font.size = Pt(11)
    lineas = md.splitlines()
    i = 0
    while i < len(lineas):
        linea = lineas[i].rstrip()
        if linea.startswith("|"):
            bloque = []
            while i < len(lineas) and lineas[i].strip().startswith("|"):
                bloque.append(lineas[i])
                i += 1
            _tabla(doc, bloque)
            continue
        m = re.match(r"^(#{1,4})\s+(.*)", linea)
        if m:
            doc.add_heading(m.group(2), level=len(m.group(1)) - 1)
        elif re.match(r"^\s*[-*]\s+", linea):
            _texto(doc.add_paragraph(style="List Bullet"), re.sub(r"^\s*[-*]\s+", "", linea))
        elif re.match(r"^\s*\d+[.)]\s+", linea):
            _texto(doc.add_paragraph(style="List Number"), re.sub(r"^\s*\d+[.)]\s+", "", linea))
        elif linea.strip():
            _texto(doc.add_paragraph(), linea.strip())
        i += 1
    doc.save(destino)


def exportar_carpeta(carpeta):
    """Convierte cada .md de <oferta>/borradores/ a .docx en <oferta>/para_subir/."""
    origen, destino = carpeta / "borradores", carpeta / "para_subir"
    destino.mkdir(exist_ok=True)
    hechos = []
    for md in sorted(origen.glob("*.md")):
        salida = destino / (md.stem + ".docx")
        md_a_docx(md.read_text(encoding="utf-8"), salida)
        hechos.append(salida)
    return hechos
