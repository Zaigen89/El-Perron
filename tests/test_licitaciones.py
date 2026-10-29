from datetime import datetime, timedelta

import pytest
import yaml

from licitaciones import api as api_mod
from licitaciones import bases, config, exportar, informes
from licitaciones.db import BD
from licitaciones.filtro import Filtro

HOY = datetime(2026, 11, 2, 10, 0)


@pytest.fixture
def filtro():
    return Filtro(config.cargar_perfil())


def detalle(nombre, descripcion="", dias=10, region="Región Metropolitana de Santiago", monto=None, categoria="86132001"):
    return {
        "CodigoExterno": "1234-56-LE26",
        "Nombre": nombre,
        "Descripcion": descripcion,
        "FechaCierre": (HOY + timedelta(days=dias)).isoformat(),
        "MontoEstimado": monto,
        "Tipo": "LE",
        "Comprador": {"NombreOrganismo": "Municipalidad de Prueba", "RegionUnidad": region},
        "Items": {"Listado": [{"CodigoCategoria": categoria, "NombreProducto": "Servicios de capacitación",
                               "Descripcion": "Taller para funcionarios", "Cantidad": 1, "UnidadMedida": "Unidad"}]},
    }


def test_perfil_y_empresa_son_yaml_validos():
    assert config.cargar_perfil()["formato"]
    assert yaml.safe_load((config.CONFIG / "empresa.yaml").read_text(encoding="utf-8"))["nombre_fantasia"]


def test_calza_capacitacion_de_liderazgo(filtro):
    puntos, motivos, descarte = filtro.evaluar(
        detalle("Capacitación en Liderazgo y Trabajo en Equipo para jefaturas"), hoy=HOY)
    assert descarte is None and puntos >= 10
    assert any("lider" in m for m in motivos)


def test_descarta_tecnicas_cerradas_y_sin_tema(filtro):
    assert filtro.evaluar(detalle("Curso de Excel avanzado"), hoy=HOY)[2].startswith("excluida")
    assert filtro.evaluar(detalle("Taller de liderazgo", dias=1), hoy=HOY)[2].startswith("cierra en")
    assert filtro.evaluar(detalle("Taller de liderazgo", dias=-1), hoy=HOY)[2] == "ya cerró"
    assert filtro.evaluar(detalle("Adquisición de resmas de papel", categoria="44121600"), hoy=HOY)[2]
    assert filtro.evaluar(detalle("Curso de soldadura"), hoy=HOY)[2]


def test_tildes_y_mayusculas_no_importan(filtro):
    assert filtro.interesa_por_nombre("JORNADA DE AUTOCUIDADO Y SALUD MENTAL")
    assert filtro.interesa_por_nombre("Capacitacion Ley Karin")
    assert not filtro.interesa_por_nombre("Compra de neumáticos")
    assert not filtro.interesa_por_nombre("Servicio de coffee break para jornada de liderazgo")


def test_region_y_monto():
    f = Filtro({**config.cargar_perfil(), "regiones": ["valparaiso"], "monto_maximo": 5_000_000})
    assert "región" in f.evaluar(detalle("Taller de liderazgo"), hoy=HOY)[2]
    assert f.evaluar(detalle("Taller de liderazgo", region="Región de Valparaíso"), hoy=HOY)[2] is None
    assert "máximo" in f.evaluar(detalle("Taller de liderazgo", region="Valparaíso", monto=9_000_000), hoy=HOY)[2]


class Respuesta:
    def __init__(self, datos, status=200):
        self.datos, self.status_code, self.text = datos, status, str(datos)

    def json(self):
        return self.datos


class SesionFalsa:
    def __init__(self, respuestas):
        self.respuestas, self.llamadas = list(respuestas), []

    def get(self, url, params, timeout):
        self.llamadas.append(params)
        return self.respuestas.pop(0)


def test_api_reintenta_peticiones_simultaneas(monkeypatch):
    monkeypatch.setattr(api_mod.time, "sleep", lambda s: None)
    sesion = SesionFalsa([
        Respuesta({"Codigo": 10500, "Mensaje": "Lo sentimos. Hemos detectado que existen peticiones simultáneas."}),
        Respuesta({"Cantidad": 1, "Listado": [{"CodigoExterno": "1-2-LE26", "Nombre": "Taller"}]}),
    ])
    cliente = api_mod.MercadoPublico("T", session=sesion)
    assert cliente.activas()[0]["CodigoExterno"] == "1-2-LE26"
    assert sesion.llamadas[0] == {"estado": "activas", "ticket": "T"}


def test_api_error_de_ticket_no_reintenta(monkeypatch):
    monkeypatch.setattr(api_mod.time, "sleep", lambda s: None)
    cliente = api_mod.MercadoPublico("X", session=SesionFalsa([Respuesta({"Codigo": 203, "Mensaje": "Ticket no válido."})]))
    with pytest.raises(api_mod.ErrorAPI, match="Ticket"):
        cliente.activas()


def test_bd_y_resumen(tmp_path):
    base = BD(tmp_path / "x.db")
    assert base.registrar_activas([{"CodigoExterno": "1234-56-LE26", "Nombre": "Taller", "FechaCierre": None}]) == 1
    assert base.registrar_activas([{"CodigoExterno": "1234-56-LE26", "Nombre": "Taller", "FechaCierre": None}]) == 0
    det = detalle("Taller de liderazgo")
    base.guardar_detalle("1234-56-LE26", det)
    base.guardar_evaluacion("1234-56-LE26", 12, ["temas: lider"], None)
    texto = informes.resumen([(base.obtener("1234-56-LE26"), base.detalle("1234-56-LE26"))], hoy=HOY)
    assert "12 pts" in texto and "1234-56-LE26" in texto and "Municipalidad de Prueba" in texto
    assert "Taller para funcionarios" in informes.ficha(det)


def test_recorte_de_bases(tmp_path):
    relleno = "\n\n".join(f"Párrafo legal genérico número {i} sin nada relevante para la oferta." for i in range(300))
    texto = relleno + "\n\nCRITERIOS DE EVALUACIÓN\nExperiencia del oferente 40%, metodología 30%, precio 30%.\n\n" + relleno
    (tmp_path / "bases").mkdir()
    (tmp_path / "bases" / "bases.txt").write_text(texto, encoding="utf-8")
    estado = bases.procesar(tmp_path)
    resumen = (tmp_path / "bases_resumen.md").read_text(encoding="utf-8")
    assert "metodología 30%" in resumen
    assert estado["caracteres_resumen"] < estado["caracteres_total"] / 10


def test_exportar_a_word(tmp_path):
    (tmp_path / "borradores").mkdir()
    (tmp_path / "borradores" / "propuesta_tecnica.md").write_text(
        "# Propuesta\n\n## 1. Metodología\nTexto con **negrita**.\n- punto\n1. paso\n\n| A | B |\n|---|---|\n| 1 | 2 |\n",
        encoding="utf-8")
    hechos = exportar.exportar_carpeta(tmp_path)
    from docx import Document
    doc = Document(str(hechos[0]))
    assert any("negrita" in p.text for p in doc.paragraphs)
    assert doc.tables[0].cell(1, 1).text == "2"
